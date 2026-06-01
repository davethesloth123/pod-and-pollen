#!/usr/bin/env python3
"""Generate the Pod & Pollen seed-data spreadsheet from the sample dataset.
Doubles as the template the Import feature will read."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ACCENT = "5F7A52"

irises = [
    # Name, Record Type, Classification, Colour Type, Status, Location, Bed / Grid Ref, Source, Planted, First Flower, Colour (summary), Height (cm), Season, Fragrance, Pod Parent, Pollen Parent, Cross Code, Generation, Favourite, Standards, Falls, Beard, Style Arms, Notes
    ["Dusky Challenger","Variety","Tall Bearded","Self","Flowering","Long Border","Row B · 3","Schreiner's, 2019","Oct 2019","May 2021","Deep blue-violet self, near-black falls",99,"Mid","Light","Navy Strut","Titan's Glory","","","Yes","Deep blue-violet, smooth velvety substance","Near-black violet, slight blue undertone on flare","Yellow at throat, tipped violet","Dark violet, crested midrib","18 May 2026 — Flowering: Three stems open, exceptional substance this year. | 2 May 2026 — Health: No sign of rot after the wet spring."],
    ["Edith Wolford","Variety","Tall Bearded","Bicolour","Flowering","Top Bed","Row A · 1","Gift from R. Hale, 2020","Sep 2020","May 2022","Yellow standards, blue-violet falls (Dykes Medal 1993)",94,"Early–Mid","Moderate","Unknown","Unknown","","","Yes","Soft butter-yellow, lightly ruffled","Rich blue-violet with darker veining at the throat","Tangerine-yellow","Pale yellow, flushed lilac at the tip","16 May 2026 — Flowering: Reliable as ever. Strong branching, 9 buds on lead stem."],
    ["Beverly Sills","Variety","Tall Bearded","Self","Growing","Top Bed","Row A · 4","Woottens, 2021","Sep 2021","May 2023","Coral-pink self, tangerine beard",89,"Mid","Sweet","Unknown","Unknown","","","","","","","","10 May 2026 — Health: Buds forming, slightly behind the others. Healthy fans."],
    ["Superstition","Variety","Tall Bearded","Self","Growing","Long Border","Row B · 7","Schreiner's, 2018","Oct 2018","May 2020","Ebony burgundy-black self",91,"Mid–Late","None","Unknown","Unknown","","","","","","","","8 May 2026 — General: Used as pollen parent for GI-26 crosses this season."],
    ["Jesse's Song","Variety","Tall Bearded","Plicata","Flowering","Top Bed","Row A · 6","Division from Dad, 2020","Sep 2020","May 2022","White ground, violet plicata margins",90,"Mid","Light","Unknown","Unknown","","","Yes","","","","","17 May 2026 — Flowering: Plicata marking very crisp this year."],
    ["Immortality","Variety","Tall Bearded (reblooming)","Self","Growing","Trial Bed","Row C · 2","Woottens, 2022","Sep 2022","May 2023","Pure white self, reliable rebloomer",76,"Mid + Autumn","Light","Unknown","Unknown","","","","","","","","6 May 2026 — General: Watching for autumn rebloom — fed in April."],
    ["GI-23-04 A","Seedling","Seedling (TB)","","First flower","Trial Bed","Row C · 9","Own cross","Apr 2024","May 2026","Rich grape-purple, gold beard — promising form",86,"Mid","Moderate","Dusky Challenger","Edith Wolford","GI-23-04","F1","Yes","","","","","22 May 2026 — Evaluation: FIRST FLOWER. Excellent substance, ruffled falls, gold beard. A keeper."],
    ["GI-23-04 C","Seedling","Seedling (TB)","","Growing","Trial Bed","Row C · 11","Own cross","Apr 2024","","Not yet flowered","","","","Dusky Challenger","Edith Wolford","GI-23-04","F1","","","","","","12 May 2026 — General: Strong fans, no buds yet — likely first flower next season."],
    ["GI-23-04 B","Seedling","Seedling (TB)","","Growing","Trial Bed","Row C · 10","Own cross","Apr 2024","","Not yet flowered — fans show yellow-tinted bases","","","","Dusky Challenger","Edith Wolford","GI-23-04","F1","","","","","","15 May 2026 — General: Vigorous — widest fans of the batch."],
    ["GI-23-04 D","Seedling","Seedling (TB)","","Watch","Holding Area","Pot 14","Own cross","Apr 2024","","Not yet flowered","","","","Dusky Challenger","Edith Wolford","GI-23-04","F1","","","","","","4 May 2026 — Health: Slower than siblings. Moved to Holding Area to assess."],
    ["GI-22-01 A","Seedling","Seedling (TB)","","First flower","Trial Bed","Row D · 1","Own cross","Apr 2023","May 2024","Warm apricot self, deep copper beard — ruffled falls",92,"Mid","Light sweet","Edith Wolford","Beverly Sills","GI-22-01","F1","Yes","Warm apricot, gently ruffled edges","Apricot with deeper copper haft, smooth substance","Copper-orange, dense","Apricot, tinged pink at the cusp","20 May 2026 — Evaluation: Third flowering — settled. Considering for naming (\"Sunday Bells\")."],
    ["GI-23-07 A","Seedling","Seedling (TB)","","Growing","Trial Bed","Row C · 14","Own cross","Apr 2024","","Not yet flowered — fans show purple stippling at base","","","","Jesse's Song","Superstition","GI-23-07","F1","","","","","","14 May 2026 — General: Strongest fan of the X2 siblings."],
]
irises_headers = ["Name","Record Type","Classification","Colour Type","Status","Location","Bed / Grid Ref","Source","Planted","First Flower","Colour (summary)","Height (cm)","Season","Fragrance","Pod Parent","Pollen Parent","Cross Code","Generation","Favourite","Standards Colour","Falls Colour","Beard Colour","Style Arms Colour","Notes"]

locations = [
    ["Top Bed","","Bed","Full sun","Sandy loam"],
    ["Long Border","","Border","Full sun","Clay loam"],
    ["Trial Bed","","Trial area","Full sun","Improved"],
    ["Greenhouse","","Greenhouse","Glass","Pots"],
    ["Holding Area","","Holding","Part shade","Pots"],
    ["Pots — Patio","Patio","Pots","Mixed","Pots"],
]
locations_headers = ["Name","Short Name","Kind","Sun","Soil"]

crosses = [
    ["GI-23-04","2023","Dusky Challenger","Edith Wolford","24 May 2023","04","Dark purple · strong branching","Evaluating","Aiming for dark purple with strong branching. Pollen kept dry 2 days."],
    ["GI-23-07","2023","Jesse's Song","Superstition","26 May 2023","07","Plicata pattern on dark ground","Growing on","Plicata × black. Hoping for striking pattern."],
    ["GI-24-02","2024","Beverly Sills","Edith Wolford","18 May 2024","02","Pink with blue infusion","Sown","Pink × yellow-blue bicolour. Long shot."],
    ["GI-24-05","2024","Dusky Challenger","Superstition","22 May 2024","05","Near-black self","Sown","Two darks together. Looking for depth and substance."],
    ["GI-22-01","2022","Edith Wolford","Beverly Sills","18 May 2022","01","Ruffled apricot bicolour","Archived","Tried for ruffled yellow-pink. Only one seedling worth keeping."],
]
crosses_headers = ["Cross Code","Season","Pod Parent","Pollen Parent","Date","Pod No.","Goal","Status","Notes"]

seed_batches = [
    ["GI-23-04","2 Aug 2023",31,"Dried, then refrigerated 8 wks","14 Oct 2023","9 Feb 2024",19,61,"20 Mar 2024","3 Apr 2024",4,0],
    ["GI-23-07","7 Aug 2023",18,"Dried, then refrigerated 8 wks","14 Oct 2023","14 Feb 2024",11,61,"20 Mar 2024","3 Apr 2024",6,0],
    ["GI-24-02","4 Aug 2024",8,"Dried, refrigerated 8 wks","12 Oct 2024","pending",0,0,"","",0,0],
    ["GI-24-05","11 Aug 2024",24,"Dried, refrigerated 8 wks","12 Oct 2024","pending",0,0,"","",0,0],
    ["GI-22-01","4 Aug 2022",22,"Dried, refrigerated 9 wks","10 Oct 2022","5 Feb 2023",14,64,"18 Mar 2023","2 Apr 2023",1,0],
]
seed_batches_headers = ["Cross Code","Harvest Date","Seeds","Treatment","Sown Date","Germ Date","Germinated","Germ %","Repot Date","Planted Out","Retained","Named"]

def add_sheet(wb, title, headers, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    header_font = Font(bold=True, color="FFFFFF", size=11)
    fill = PatternFill("solid", fgColor=ACCENT)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=False)
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            ws.cell(row=r, column=c, value=val)
    # column widths
    for c, h in enumerate(headers, 1):
        maxlen = max([len(str(h))] + [len(str(row[c-1])) for row in rows if c-1 < len(row)])
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = min(max(maxlen + 2, 12), 60)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

wb = openpyxl.Workbook()
add_sheet(wb, "Irises", irises_headers, irises, first=True)
add_sheet(wb, "Locations", locations_headers, locations)
add_sheet(wb, "Crosses", crosses_headers, crosses)
add_sheet(wb, "Seed Batches", seed_batches_headers, seed_batches)

out = "/home/claude/repo/seed/pod_and_pollen_seed_data.xlsx"
wb.save(out)
print("Saved", out)
print("Sheets:", wb.sheetnames)
